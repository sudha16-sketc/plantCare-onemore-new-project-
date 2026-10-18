"""
PPT visual guide generator.

Generates a real PowerPoint (.pptx) file with python-pptx and returns a
sanitized filename. Downloads are served through the controlled
/api/v1/files/{filename} endpoint (never a raw static mount that would leak
unrelated files), and stale files are cleaned up on a TTL.
"""

import asyncio
import logging
import os
import re
import time
import uuid
from datetime import datetime, timezone

from pptx import Presentation

from app.schemas.plant import GeminiResponse
from app.utils.config import settings

logger = logging.getLogger("plantcare.ppt")

_SLUG_RE = re.compile(r"[^a-z0-9]+")


def slugify(value: str, max_len: int = 60) -> str:
    """Sanitize a plant name into a safe filename slug (no path separators)."""
    slug = _SLUG_RE.sub("_", value.lower()).strip("_")
    return slug[:max_len] or "plant"


class PPTService:
    """Generates PPTX care guides from structured GeminiResponse data."""

    def __init__(self) -> None:
        self.output_dir = settings.generated_dir
        os.makedirs(self.output_dir, exist_ok=True)

    @property
    def configured(self) -> bool:
        return True  # python-pptx is local; generation is always available

    async def generate(self, plant_care_data: GeminiResponse, plant_name: str) -> dict:
        """
        Generate a PowerPoint presentation from plant care data.

        Returns {"file_name": ..., "file_url": ...} where file_url is a
        relative API path that clients must resolve against the API origin
        (the route layer converts it to an absolute URL using the request).
        """
        # Unique suffix avoids collisions between same-named plants and
        # defeats any residual path trickery in the plant name.
        file_name = f"{slugify(plant_name)}_care_guide_{uuid.uuid4().hex[:8]}.pptx"
        file_path = os.path.join(self.output_dir, file_name)

        # Heavy work runs in a thread so the event loop stays responsive.
        await asyncio.to_thread(self._build_presentation, plant_care_data, plant_name, file_path)

        return {
            "status": "success",
            "file_name": file_name,
            "file_url": f"/api/v1/files/{file_name}",
            "file_type": "pptx",
            "message": "Care guide presentation generated successfully",
        }

    def _build_presentation(self, data: GeminiResponse, plant_name: str, file_path: str) -> None:
        prs = Presentation()

        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = f"{plant_name} Care Guide"
        slide.placeholders[1].text = f"Difficulty: {data.plant_overview.difficulty_level}"

        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Plant Overview"
        slide.placeholders[1].text = data.plant_overview.description

        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Ideal Growing Conditions"
        conditions = data.plant_overview.ideal_conditions
        cond = conditions if isinstance(conditions, dict) else {}
        slide.placeholders[1].text = (
            f"Temperature: {cond.get('temperature', 'Not specified')}\n"
            f"Humidity: {cond.get('humidity', 'Not specified')}\n"
            f"Sunlight: {cond.get('sunlight', 'Not specified')}\n"
            f"Soil pH: {cond.get('soil_ph', 'Not specified')}"
        )

        for stage in data.growth_stages:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = stage.stage_name
            slide.placeholders[1].text = (
                f"Duration: {stage.duration}\n\n"
                f"Care:\n{stage.care_instructions}\n\n"
                "Indicators:\n- " + "\n- ".join(stage.key_indicators)
            )

        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Daily Care Routine"
        daily = data.daily_care
        slide.placeholders[1].text = (
            "Morning:\n- " + "\n- ".join(daily.morning_routine) + "\n\n"
            "Afternoon:\n- " + "\n- ".join(daily.afternoon_routine) + "\n\n"
            "Evening:\n- " + "\n- ".join(daily.evening_routine)
        )

        if daily.weekly_tasks:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = "Weekly Tasks"
            slide.placeholders[1].text = "- " + "\n- ".join(daily.weekly_tasks)

        for problem in data.common_problems:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = problem.problem
            slide.placeholders[1].text = (
                "Symptoms:\n- " + "\n- ".join(problem.symptoms) + "\n\n"
                f"Solution:\n{problem.solution}\n\n"
                f"Prevention:\n{problem.prevention}"
            )

        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Additional Tips"
        slide.placeholders[1].text = "- " + "\n- ".join(data.additional_tips)

        prs.save(file_path)

    def cleanup_expired(self) -> int:
        """Delete generated files older than the configured TTL. Returns count removed."""
        cutoff = time.time() - settings.generated_file_ttl_seconds
        removed = 0
        try:
            for name in os.listdir(self.output_dir):
                path = os.path.join(self.output_dir, name)
                if not os.path.isfile(path):
                    continue
                if not name.endswith(".pptx"):
                    continue
                try:
                    if os.path.getmtime(path) < cutoff:
                        os.remove(path)
                        removed += 1
                except OSError:
                    continue
        except OSError as exc:
            logger.warning("Generated-file cleanup failed: %s", exc)
        return removed

    def safe_path(self, file_name: str):
        """
        Resolve a requested file name to a path inside output_dir, or None.

        Rejects anything that is not a plain generated .pptx filename
        (path separators, traversal, unexpected extensions).
        """
        if not file_name or "/" in file_name or "\\" in file_name or ".." in file_name:
            return None
        if not re.fullmatch(r"[a-z0-9_\-]{1,120}\.pptx", file_name):
            return None
        path = os.path.abspath(os.path.join(self.output_dir, file_name))
        base = os.path.abspath(self.output_dir)
        if not path.startswith(base + os.sep):
            return None
        return path if os.path.isfile(path) else None

    async def start_cleanup_task(self) -> "asyncio.Task":
        """Background loop that periodically removes expired files."""

        async def _loop() -> None:
            while True:
                await asyncio.sleep(max(60, settings.generated_file_ttl_seconds // 4))
                removed = await asyncio.to_thread(self.cleanup_expired)
                if removed:
                    logger.info("Cleaned up %d expired generated file(s)", removed)

        return asyncio.create_task(_loop())


ppt_service = PPTService()


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()
